# %% [markdown]
# ## Setup
#
# Builds one long lookup table combining precinct-level rent figures (as
# published) and region-level aggregates (computed here, count-weighted
# average of precinct medians) — the two granularities the suburb→precinct
# mapping actually needs, in one place.

# %%
import openpyxl
import pandas as pd

RENT_XLSX = "data/moving_annual_rent_by_suburb.xlsx"
OUTPUT_CSV = "data/rent_benchmarks_long.csv"

# %% [markdown]
# ## Parse every sheet into a long precinct-level table
#
# Each sheet has the same layout: row 2 = quarter labels (repeated across a
# Count/Median pair), row 3 = "Count"/"Median" labels, row 4+ = data with
# region filled only on a precinct's first row (forward-fill needed).
# Suppressed cells are the literal string "-", not blank — those must be
# excluded, not treated as zero.

# %%
wb = openpyxl.load_workbook(RENT_XLSX, read_only=True, data_only=True)

precinct_rows = []

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    all_rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    quarter_header = all_rows[1]
    count_median_header = all_rows[2]

    quarter_pairs = []  # [(quarter_label, count_col_idx, median_col_idx), ...]
    col = 2
    while col + 1 < len(quarter_header):
        label = quarter_header[col]
        if label and count_median_header[col] == "Count" and count_median_header[col + 1] == "Median":
            quarter_pairs.append((label, col, col + 1))
        col += 2

    current_region = None
    for row in all_rows[3:]:
        region, precinct = row[0], row[1]
        if region:
            current_region = region
        if not precinct:
            continue
        for quarter_label, count_idx, median_idx in quarter_pairs:
            count, median = row[count_idx], row[median_idx]
            if count in (None, "-") or median in (None, "-"):
                continue
            precinct_rows.append({
                "dwelling_sheet": sheet_name,
                "region": current_region,
                "precinct": precinct,
                "quarter": quarter_label,
                "count": count,
                "median_weekly_rent": median,
            })

precinct_df = pd.DataFrame(precinct_rows)
print("Precinct-level rows:", len(precinct_df))
print(precinct_df.head())

# %% [markdown]
# ## Region-level aggregation
#
# Count-weighted average of precinct medians per (sheet, region, quarter).
# This is an approximation of the true region median (it averages medians,
# not raw bond lodgements — the raw lodgement data isn't published) but it's
# the standard practical way to roll up this kind of report, and it's
# weighted so a 200-lodgement precinct counts more than a 10-lodgement one.

# %%
def weighted_median(group: pd.DataFrame) -> pd.Series:
    total_count = group["count"].sum()
    weighted = (group["count"] * group["median_weekly_rent"]).sum() / total_count
    return pd.Series({
        "median_weekly_rent": weighted,
        "count": total_count,
        "n_precincts": len(group),
    })


region_df = (
    precinct_df.groupby(["dwelling_sheet", "region", "quarter"])
    .apply(weighted_median, include_groups=False)
    .reset_index()
)
region_df["precinct"] = None
print("Region-level rows:", len(region_df))
print(region_df.head())

# %% [markdown]
# ## Combine into one long lookup table

# %%
precinct_df["granularity"] = "precinct"
region_df["granularity"] = "region"

benchmarks = pd.concat([precinct_df, region_df], ignore_index=True)
benchmarks = benchmarks[
    ["granularity", "dwelling_sheet", "region", "precinct", "quarter", "median_weekly_rent", "count", "n_precincts"]
]
benchmarks.loc[benchmarks["granularity"] == "precinct", "n_precincts"] = 1

benchmarks.to_csv(OUTPUT_CSV, index=False)
print(f"Saved {len(benchmarks)} rows to {OUTPUT_CSV}")

# %% [markdown]
# ## Sanity checks
#
# The property dataset's sale dates run 2016-2018 — confirm that window
# actually has data (not just suppressed/missing), and check how many
# region/quarter combos ended up with zero contributing precincts (should
# be none, since region rows only exist where at least one precinct fed in).

# %%
sale_period_quarters = [f"{q} {y}" for y in (2016, 2017, 2018) for q in ("Mar", "Jun", "Sep", "Dec")]
coverage = benchmarks[benchmarks["quarter"].isin(sale_period_quarters)]
print("Rows covering the 2016-2018 sale window:", len(coverage))
print(
    coverage.groupby(["dwelling_sheet", "granularity"])["quarter"].nunique()
    .rename("quarters_present")
    .reset_index()
)

zero_precinct_regions = region_df[region_df["n_precincts"] == 0]
print(f"\nRegion/quarter combos with zero contributing precincts: {len(zero_precinct_regions)}")
